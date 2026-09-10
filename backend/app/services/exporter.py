"""XLSX export for BYTE2DNA-POC-1 sequences."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.constants import APPLICATION_NAME, ENCODING_VERSION, PROFILE_NAME
from app.models import EncodingResult, RoundtripResult


HEADER_FILL = PatternFill("solid", fgColor="0F766E")
HEADER_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FEF3C7")
FAIL_FILL = PatternFill("solid", fgColor="FEE2E2")
PASS_FILL = PatternFill("solid", fgColor="DCFCE7")


def _style_header(sheet: Worksheet, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = sheet.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(columns)}1"


def _autosize(sheet: Worksheet, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[get_column_letter(col)].width = width


def _write_kv(sheet: Worksheet, rows: list[tuple[str, object]]) -> None:
    sheet.append(["Key", "Value"])
    _style_header(sheet, 2)
    for key, value in rows:
        sheet.append([key, value if value is not None else ""])
    _autosize(sheet, {1: 36, 2: 88})


def _sequences_sheet(wb: Workbook, result: EncodingResult) -> None:
    sheet = wb.active
    sheet.title = "SEQUENCES"
    headers = [
        "Fragment_ID",
        "Fragment_Index",
        "DNA_Sequence",
        "Length_nt",
        "GC_Percent",
        "Max_Homopolymer",
        "Encoding_Variant",
        "ECC",
        "Status",
    ]
    sheet.append(headers)
    _style_header(sheet, len(headers))
    for fragment in result.fragments:
        sheet.append(
            [
                fragment.fragment_id,
                fragment.index + 1,
                fragment.sequence,
                fragment.length_nt,
                fragment.gc_percent,
                fragment.max_homopolymer,
                fragment.encoding_variant,
                "enabled" if fragment.ecc else "disabled",
                fragment.status,
            ]
        )
        status_cell = sheet.cell(sheet.max_row, 9)
        if fragment.status == "WARNING":
            status_cell.fill = WARN_FILL
        elif fragment.status == "INVALID":
            status_cell.fill = FAIL_FILL
        else:
            status_cell.fill = PASS_FILL
    _autosize(sheet, {1: 14, 2: 16, 3: 80, 4: 12, 5: 14, 6: 18, 7: 18, 8: 12, 9: 12})


def _metadata_sheet(
    wb: Workbook,
    result: EncodingResult,
    roundtrip: RoundtripResult,
) -> None:
    sheet = wb.create_sheet("METADATA")
    file_info = result.file
    stats = result.stats
    config = result.config
    sha_export = (
        file_info.sha256_original if config.include_sha256_export else "(omitted from export)"
    )
    sha_recon = (
        roundtrip.sha256_reconstructed
        if config.include_sha256_export
        else "(omitted from export)"
    )
    _write_kv(
        sheet,
        [
            ("Application", APPLICATION_NAME),
            ("Encoding_Version", ENCODING_VERSION),
            ("Profile", PROFILE_NAME),
            ("Original_Filename", file_info.original_filename),
            ("Original_Extension", file_info.original_extension),
            ("Original_Size", file_info.original_size),
            ("SHA256_Original", sha_export),
            ("Compression", "ON" if config.compression else "OFF"),
            ("Compressed_Size", stats.compressed_size),
            ("Fragment_Length", config.fragment_length),
            ("Fragment_Count", stats.fragment_count),
            ("GC_Min_Config", config.gc_min),
            ("GC_Max_Config", config.gc_max),
            ("Homopolymer_Max_Config", config.homopolymer_max),
            ("ECC", "ON" if config.ecc else "OFF"),
            ("Total_Nucleotides", stats.total_nucleotides),
            ("Roundtrip_Result", roundtrip.result),
            ("SHA256_Reconstructed", sha_recon),
        ],
    )


def _status_for_metric(ok: bool) -> str:
    return "PASS" if ok else "FAIL"


def _qc_sheet(wb: Workbook, result: EncodingResult, roundtrip: RoundtripResult) -> None:
    sheet = wb.create_sheet("QC")
    stats = result.stats
    config = result.config
    gc_ok = config.gc_min <= stats.gc_min_observed and stats.gc_max_observed <= config.gc_max
    hp_ok = stats.max_homopolymer_observed <= config.homopolymer_max
    rows = [
        ("Metric", "Value", "Status"),
        ("GC average", stats.gc_average, _status_for_metric(gc_ok)),
        ("GC minimum", stats.gc_min_observed, _status_for_metric(stats.gc_min_observed >= config.gc_min)),
        ("GC maximum", stats.gc_max_observed, _status_for_metric(stats.gc_max_observed <= config.gc_max)),
        (
            "Maximum homopolymer",
            stats.max_homopolymer_observed,
            _status_for_metric(hp_ok),
        ),
        ("Fragments total", stats.fragment_count, "INFO"),
        ("Fragments valid", stats.fragments_valid, "INFO"),
        ("Fragments warning", stats.fragments_warning, "WARN" if stats.fragments_warning else "PASS"),
        (
            "Fragments invalid",
            stats.fragments_invalid,
            "FAIL" if stats.fragments_invalid else "PASS",
        ),
        ("Round-trip", roundtrip.result, roundtrip.result),
    ]
    for row in rows:
        sheet.append(list(row))
    _style_header(sheet, 3)
    _autosize(sheet, {1: 28, 2: 22, 3: 12})

    # Simple GC histogram buckets for optional chart.
    hist_start = len(rows) + 3
    sheet.cell(hist_start, 1, "GC_Bucket")
    sheet.cell(hist_start, 2, "Count")
    buckets = list(range(0, 100, 10))
    counts = [0] * len(buckets)
    for fragment in result.fragments:
        bucket = min(int(fragment.gc_percent // 10), 9)
        counts[bucket] += 1
    for i, start in enumerate(buckets):
        sheet.cell(hist_start + 1 + i, 1, f"{start}-{start + 10}")
        sheet.cell(hist_start + 1 + i, 2, counts[i])
    chart = BarChart()
    chart.title = "GC percent histogram"
    chart.y_axis.title = "Fragments"
    data = Reference(sheet, min_col=2, min_row=hist_start, max_row=hist_start + 10)
    cats = Reference(sheet, min_col=1, min_row=hist_start + 1, max_row=hist_start + 10)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    sheet.add_chart(chart, "E2")


def _decoding_sheet(wb: Workbook, result: EncodingResult) -> None:
    sheet = wb.create_sheet("DECODING")
    meta = result.decode_metadata
    mapping = ", ".join(f"{bits}→{base}" for base, bits in sorted(
        {v: k for k, v in meta.base_mapping.items()}.items()
    )) if False else "00→A ; 01→C ; 10→G ; 11→T"
    _write_kv(
        sheet,
        [
            ("Encoding version", ENCODING_VERSION),
            ("Base mapping", mapping),
            (
                "Fragment structure",
                "[variant_header 8 nt][scrambled body][filler to target length]",
            ),
            (
                "Index encoding",
                "32-bit big-endian index stored in the first 4 descrambled body bytes",
            ),
            (
                "Variant/seed mechanism",
                "8-nt GC-balanced header encodes variant 0-255; body is XOR-scrambled "
                "with a SHA-256 keystream keyed by BYTE2DNA-POC-1 + variant. Encoder "
                "tries variants until GC/homopolymer constraints pass, else best effort INVALID.",
            ),
            (
                "Padding rule",
                "Last payload chunk is zero-padded to payload_capacity. Inner header "
                "stores original_size and data_size so padding is discarded on decode. "
                "Leftover nucleotides are ACGT filler and ignored.",
            ),
            (
                "Inner payload",
                "magic B2D1 | flags (bit0=zlib) | original_size uint32 BE | "
                "data_size uint32 BE | payload bytes",
            ),
            ("Compression", "zlib" if result.config.compression else "none"),
            (
                "ECC configuration",
                f"reedsolo RSCodec nsym={meta.ecc_nsym}" if result.config.ecc else "disabled",
            ),
            ("Payload capacity (bytes/fragment)", meta.payload_capacity),
            ("Fragment length (nt)", meta.fragment_length),
            (
                "Reconstruction",
                "1. Read variant header. 2. Convert body DNA to bytes. 3. XOR descramble. "
                "4. Optional RS decode. 5. Read index + payload. 6. Sort by index. "
                "7. Concatenate. 8. Parse inner header. 9. Optional zlib decompress.",
            ),
        ],
    )


def build_xlsx(result: EncodingResult, roundtrip: RoundtripResult) -> bytes:
    wb = Workbook()
    _sequences_sheet(wb, result)
    _metadata_sheet(wb, result, roundtrip)
    _qc_sheet(wb, result, roundtrip)
    _decoding_sheet(wb, result)
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def export_filename(original_filename: str, when: datetime | None = None) -> str:
    stamp = (when or datetime.now(timezone.utc)).strftime("%Y%m%d%H%M%S")
    stem = Path(original_filename).stem or "file"
    safe = "".join(ch if ch.isalnum() or ch in "-_" else "-" for ch in stem)[:80]
    return f"byte2dna_{safe}_{stamp}.xlsx"
