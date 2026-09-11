"""Extract DNA sequences and encoding metadata from XLSX and XLS workbooks."""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import openpyxl

from app.constants import (
    APPLICATION_NAME,
    ECC_NSYM,
    ENCODING_VERSION,
    INDEX_BYTES,
    VARIANT_HEADER_NT,
)
from app.models import DecodeMetadata
from app.services.dna import sha256_hex
from app.services.fragmenter import payload_capacity

try:
    import xlrd  # type: ignore[import-untyped]
except ImportError:
    xlrd = None  # type: ignore[assignment]


class ExcelParseError(ValueError):
    """Raised when an Excel workbook cannot be parsed for DNA sequences."""


def _parse_bool_value(val: Any, default: bool = True) -> bool:
    if val is None:
        return default
    if isinstance(val, bool):
        return val
    s = str(val).strip().upper()
    if s in {"ON", "TRUE", "1", "YES", "VRAI", "OUI"}:
        return True
    if s in {"OFF", "FALSE", "0", "NO", "FAUX", "NON"}:
        return False
    return default


def _read_xlsx(content: bytes) -> tuple[list[str], dict[str, Any]]:
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ExcelParseError("Invalid XLSX file") from exc

    metadata: dict[str, Any] = {}
    if "METADATA" in wb.sheetnames:
        meta_sheet = wb["METADATA"]
        for row in meta_sheet.iter_rows(values_only=True):
            if row and len(row) >= 2 and row[0] is not None:
                metadata[str(row[0]).strip()] = row[1]

    # Look for sequences sheet
    sheet_name = "SEQUENCES" if "SEQUENCES" in wb.sheetnames else wb.sheetnames[0]
    seq_sheet = wb[sheet_name]

    sequences: list[str] = []
    seq_col_idx: int | None = None
    first_row = True

    for row in seq_sheet.iter_rows(values_only=True):
        if not row:
            continue
        if first_row:
            first_row = False
            # Try to identify sequence column by header
            for idx, cell in enumerate(row):
                header_str = str(cell or "").strip().lower()
                if "sequence" in header_str or "dna" in header_str:
                    seq_col_idx = idx
                    break
            # If not found by name, default to index 2 (column 3 in Byte2DNA export)
            if seq_col_idx is None:
                seq_col_idx = 2 if len(row) > 2 else 0
            # If this row itself looks like a DNA sequence (no header), include it
            first_val = str(row[seq_col_idx] or "").strip().upper()
            if first_val and len(first_val) >= 4 and all(c in "ACGT" for c in first_val):
                sequences.append(first_val)
            continue

        if seq_col_idx < len(row):
            val = str(row[seq_col_idx] or "").strip().upper()
            if val and all(c in "ACGT" for c in val):
                sequences.append(val)

    wb.close()
    return sequences, metadata


def _read_xls(content: bytes) -> tuple[list[str], dict[str, Any]]:
    if xlrd is None:
        raise ExcelParseError("XLS support requires xlrd library")
    try:
        book = xlrd.open_workbook(file_contents=content)
    except Exception as exc:
        raise ExcelParseError("Invalid XLS file") from exc

    metadata: dict[str, Any] = {}
    if "METADATA" in book.sheet_names():
        meta_sheet = book.sheet_by_name("METADATA")
        for row_idx in range(meta_sheet.nrows):
            key = meta_sheet.cell_value(row_idx, 0)
            val = meta_sheet.cell_value(row_idx, 1) if meta_sheet.ncols > 1 else None
            if key:
                metadata[str(key).strip()] = val

    sheet_name = "SEQUENCES" if "SEQUENCES" in book.sheet_names() else book.sheet_names()[0]
    seq_sheet = book.sheet_by_name(sheet_name)

    sequences: list[str] = []
    seq_col_idx: int | None = None
    start_row = 0

    if seq_sheet.nrows > 0:
        header_row = [seq_sheet.cell_value(0, c) for c in range(seq_sheet.ncols)]
        for idx, cell in enumerate(header_row):
            header_str = str(cell or "").strip().lower()
            if "sequence" in header_str or "dna" in header_str:
                seq_col_idx = idx
                start_row = 1
                break
        if seq_col_idx is None:
            first_val = str(header_row[2] if len(header_row) > 2 else header_row[0]).strip().upper()
            if first_val and all(c in "ACGT" for c in first_val):
                seq_col_idx = 2 if len(header_row) > 2 else 0
                start_row = 0
            else:
                seq_col_idx = 2 if len(header_row) > 2 else 0
                start_row = 1

    for row_idx in range(start_row, seq_sheet.nrows):
        if seq_col_idx < seq_sheet.ncols:
            val = str(seq_sheet.cell_value(row_idx, seq_col_idx) or "").strip().upper()
            if val and all(c in "ACGT" for c in val):
                sequences.append(val)

    return sequences, metadata


def extract_sequences_and_metadata(
    content: bytes,
    filename: str = "export.xlsx",
) -> tuple[list[str], dict[str, Any]]:
    """Extract sequences and metadata from an uploaded XLSX or XLS file."""
    if not content:
        raise ExcelParseError("Empty file")

    ext = Path(filename).suffix.lower()
    is_zip = content.startswith(b"PK")
    is_cfb = content.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")

    if ext == ".xlsx" or is_zip:
        sequences, metadata = _read_xlsx(content)
    elif ext == ".xls" or is_cfb:
        sequences, metadata = _read_xls(content)
    else:
        raise ExcelParseError("Unsupported format. Please upload an .xlsx or .xls file.")

    if not sequences:
        raise ExcelParseError("No valid DNA sequences found in the Excel workbook.")

    return sequences, metadata


def build_decode_metadata_from_dict(
    metadata: dict[str, Any],
    sequences: list[str],
) -> tuple[DecodeMetadata, str, str | None]:
    """Construct DecodeMetadata and extract target filename and sha256 from Excel metadata."""
    sample_len = len(sequences[0])
    comp_val = metadata.get("Compression", metadata.get("compression", "ON"))
    compression = _parse_bool_value(comp_val, default=True)

    ecc_val = metadata.get("ECC", metadata.get("ecc", "ON"))
    ecc = _parse_bool_value(ecc_val, default=True)

    frag_len_raw = metadata.get("Fragment_Length", metadata.get("fragment_length"))
    fragment_length = int(frag_len_raw) if frag_len_raw is not None else sample_len

    nsym = ECC_NSYM if ecc else 0
    cap = payload_capacity(fragment_length, ecc)

    orig_name = str(metadata.get("Original_Filename") or "").strip()
    if not orig_name:
        ext = str(metadata.get("Original_Extension") or "").strip().lstrip(".")
        orig_name = f"reconstructed_document.{ext}" if ext else "reconstructed_document.doc"

    sha_orig = str(metadata.get("SHA256_Original") or "").strip()
    if not sha_orig or sha_orig in {"(omitted from export)", "None", "null"}:
        sha_orig = None

    decode_metadata = DecodeMetadata(
        encoding_version=str(metadata.get("Encoding_Version") or ENCODING_VERSION),
        compression=compression,
        ecc=ecc,
        ecc_nsym=nsym,
        fragment_length=fragment_length,
        payload_capacity=cap,
        variant_header_nt=VARIANT_HEADER_NT,
        index_bytes=INDEX_BYTES,
        base_mapping={"A": "00", "C": "01", "G": "10", "T": "11"},
    )

    return decode_metadata, orig_name, sha_orig
