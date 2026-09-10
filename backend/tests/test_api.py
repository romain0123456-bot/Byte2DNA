from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app
from tests.conftest import tiny_docx_bytes, tiny_pdf_bytes

client = TestClient(app)


def test_health() -> None:
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def _encode(filename: str, content: bytes, **params: object):
    files = {"file": (filename, content, "application/octet-stream")}
    data = {
        "compression": "true",
        "include_sha256_export": "true",
        "ecc": "true",
        "fragment_length": "150",
        "gc_min": "40",
        "gc_max": "60",
        "homopolymer_max": "3",
        **{k: str(v).lower() if isinstance(v, bool) else str(v) for k, v in params.items()},
    }
    return client.post("/api/encode", files=files, data=data)


def test_pdf_and_docx_encode_roundtrip() -> None:
    pdf = tiny_pdf_bytes(b"hello-pdf-bytes")
    docx = tiny_docx_bytes(b"hello-docx-bytes")
    for name, content in (("sample.pdf", pdf), ("sample.docx", docx)):
        response = _encode(name, content)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["encoding_version"] == "BYTE2DNA-POC-1"
        assert body["roundtrip"]["result"] == "PASS"
        assert body["roundtrip"]["bit_identical"] is True
        assert body["roundtrip"]["sha256_original"] == body["roundtrip"]["sha256_reconstructed"]
        assert body["export_allowed"] is True
        assert body["file"]["original_filename"] == name
        assert len(body["file"]["sha256_original"]) == 64


def test_invalid_file_handling() -> None:
    assert _encode("notes.txt", b"hello").status_code == 400
    assert _encode("empty.pdf", b"").json()["error"] == "Empty file"
    huge = b"%PDF" + b"a" * (5 * 1024 * 1024)
    assert _encode("big.pdf", huge).json()["error"] == "File too large"


def test_upload_over_5mb_rejected_before_encoding() -> None:
    oversized = b"%PDF" + b"x" * (5 * 1024 * 1024 + 10)
    response = _encode("huge.pdf", oversized)
    assert response.status_code == 400
    data = response.json()
    assert data["error"] == "File too large"
    assert data["detail"] == "File too large"


def test_compression_toggle_via_api() -> None:
    content = tiny_pdf_bytes(b"AAAA" * 200)
    on = _encode("c.pdf", content, compression=True).json()
    off = _encode("c.pdf", content, compression=False).json()
    assert on["roundtrip"]["result"] == "PASS"
    assert off["roundtrip"]["result"] == "PASS"
    assert on["stats"]["compressed_size"] <= off["stats"]["compressed_size"]


def test_export_requires_roundtrip_and_has_sheets() -> None:
    content = tiny_pdf_bytes(b"export-me")
    encoded = _encode("document-test.pdf", content).json()
    result_id = encoded["result_id"]
    response = client.post("/api/export", json={"result_id": result_id})
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    wb = load_workbook(BytesIO(response.content))
    assert set(wb.sheetnames) == {"SEQUENCES", "METADATA", "QC", "DECODING"}

    missing = client.post("/api/export", json={"result_id": "0" * 32})
    assert missing.status_code == 404
