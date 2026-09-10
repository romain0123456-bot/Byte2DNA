from __future__ import annotations

from io import BytesIO

from openpyxl import load_workbook

from app.models import EncodeConfig
from app.services.codec import encode
from app.services.exporter import build_xlsx, export_filename
from app.services.validator import validate_roundtrip


def test_xlsx_has_required_sheets_and_columns() -> None:
    original = b"xlsx-export-payload-" + bytes(range(80))
    result = encode(original, EncodeConfig(include_sha256_export=True), filename="document-test.pdf")
    roundtrip = validate_roundtrip(original, result.fragments, result.decode_metadata)
    assert roundtrip.result == "PASS"
    payload = build_xlsx(result, roundtrip)
    wb = load_workbook(BytesIO(payload))
    assert wb.sheetnames == ["SEQUENCES", "METADATA", "QC", "DECODING"]

    seq = wb["SEQUENCES"]
    headers = [cell.value for cell in seq[1]]
    assert headers == [
        "Fragment_ID",
        "Fragment_Index",
        "DNA_Sequence",
        "Length_nt",
        "GC_Percent",
        "Max_Homopolymer",
        "Encoding_Variant",
        "ECC",
        "Status",
    ]
    first_seq = seq.cell(2, 3).value
    assert isinstance(first_seq, str) and set(first_seq) <= set("ACGT")
    assert len(first_seq) == result.config.fragment_length

    meta_keys = {row[0] for row in wb["METADATA"].iter_rows(min_row=2, values_only=True) if row[0]}
    for key in (
        "Application",
        "Encoding_Version",
        "Original_Filename",
        "Original_Extension",
        "Original_Size",
        "SHA256_Original",
        "Compression",
        "Compressed_Size",
        "Fragment_Length",
        "Fragment_Count",
        "GC_Min_Config",
        "GC_Max_Config",
        "Homopolymer_Max_Config",
        "ECC",
        "Total_Nucleotides",
        "Roundtrip_Result",
        "SHA256_Reconstructed",
    ):
        assert key in meta_keys

    qc_metrics = {row[0] for row in wb["QC"].iter_rows(min_row=2, values_only=True) if row[0]}
    for key in (
        "GC average",
        "GC minimum",
        "GC maximum",
        "Maximum homopolymer",
        "Fragments total",
        "Fragments valid",
        "Fragments warning",
        "Fragments invalid",
        "Round-trip",
    ):
        assert key in qc_metrics

    decoding_keys = {row[0] for row in wb["DECODING"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert "Encoding version" in decoding_keys
    assert "Base mapping" in decoding_keys
    assert "Fragment structure" in decoding_keys


def test_export_filename_pattern() -> None:
    name = export_filename("document-test.pdf")
    assert name.startswith("byte2dna_document-test_")
    assert name.endswith(".xlsx")
