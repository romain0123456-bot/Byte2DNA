"""Sprint 5 end-to-end probe: encode -> export XLSX -> decode through the HTTP API.

Run from repo root: PYTHONPATH=backend:backend/tests python qualification/sprint5/e2e_probe.py
Uses the synthetic fixtures of backend/tests/conftest.py. Prints observed values only.
"""
import hashlib
import io

from conftest import tiny_doc_bytes, tiny_docx_bytes, tiny_pdf_bytes
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
CASES = [
    ("fixture.pdf", tiny_pdf_bytes(), "application/pdf"),
    ("fixture.docx", tiny_docx_bytes(),
     "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    ("fixture.doc", tiny_doc_bytes(), "application/msword"),
]
ok = True
for name, data, mime in CASES:
    enc = client.post("/api/encode", files={"file": (name, data, mime)}).json()
    sha_in = hashlib.sha256(data).hexdigest()
    exp = client.get("/api/export", params={"result_id": enc["result_id"]})
    wb = load_workbook(io.BytesIO(exp.content))
    dec = client.post("/api/decode", files={"file": ("export.xlsx", exp.content, "application/octet-stream")}).json()
    out = client.get(dec["download_url"]).content
    same = out == data and hashlib.sha256(out).hexdigest() == sha_in
    ok &= same and enc["roundtrip"]["result"] == "PASS" and exp.status_code == 200
    print(f"{name}: size={len(data)} sha256={sha_in}")
    print(f"  encode roundtrip={enc['roundtrip']['result']} export_allowed={enc['export_allowed']} stats={enc['stats']}")
    print(f"  export status={exp.status_code} sheets={wb.sheetnames}")
    print(f"  decode sha256_matched={dec['sha256_matched']} decoded_sha256={dec['sha256']}")
    print(f"  downloaded bytes identical={out == data} sha256 identical={hashlib.sha256(out).hexdigest() == sha_in}")
print("ALL_OK" if ok else "FAILURE")
raise SystemExit(0 if ok else 1)
